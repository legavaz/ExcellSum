import base64
import io

import openpyxl
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app import app
from storage import STORAGE_DIR

client = TestClient(app)
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def post_convert(rows, data):
    return client.post(
        "/convert",
        files={"file": ("t.xlsx", make_xlsx(rows), XLSX_MIME)},
        data=data,
    )


def test_ui_page_served():
    r = client.get("/")
    assert r.status_code == 200
    assert "text/html" in r.headers["content-type"]
    assert 'id="form"' in r.text
    assert 'id="download"' in r.text
    assert "/convert" in r.text


def test_health():
    r = client.get("/health")
    assert r.status_code == 200
    assert r.json() == {"status": "ok"}


def test_openapi_has_no_auth():
    schema = client.get("/openapi.json").json()
    assert "securitySchemes" not in schema.get("components", {})
    assert "security" not in schema["paths"]["/convert"]["post"]


def test_convert_bottom_returns_xlsx():
    rows = [["Ур.", "Наим.", "Сумма"],
            [0, "Группа A", None],
            [1, "x", 10],
            [1, "y", 20],
            [0, "Итого A", 30]]
    r = post_convert(rows, {"columns": "C", "level_column": "A"})
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert r.headers["X-Job-Id"]
    assert r.headers["X-Log-Preview"]
    log = base64.b64decode(r.headers["X-Log-Preview-B64"]).decode("utf-8")
    assert "Итогов" in log
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws["C5"].value == "=SUM(C3:C4)"


def test_convert_grand_total_flat_report():
    rows = [["Сотрудник", "Часы"],
            ["x", 10],
            ["y", 20],
            ["Итого", 30]]
    r = post_convert(rows, {"columns": "B", "grand_total": "true"})
    assert r.status_code == 200, r.text
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws["B4"].value == "=SUM(B2:B3)"


def test_convert_dry_run_returns_json_without_output_file():
    rows = [[0, "Итого A", 30], [1, "x", 10]]
    r = post_convert(rows, {"columns": "C", "level_column": "A",
                            "dry_run": "true"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["dry_run"] is True
    assert "=" in body["log"]
    assert "Сохранено" not in body["log"]
    job_dir = STORAGE_DIR / r.headers["X-Job-Id"]
    assert list(job_dir.glob("fixed_*")) == []


def post_analyze(rows, data=None):
    return client.post(
        "/analyze",
        files={"file": ("t.xlsx", make_xlsx(rows), XLSX_MIME)},
        data=data or {},
    )


def test_analyze_flat_report():
    rows = [["Сотрудник", "Часы"],
            ["x", 10],
            ["y", 20],
            ["Итого", 30]]
    r = post_analyze(rows)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["flat"] is True
    assert body["grand_total"] is True
    assert body["total_row"] == 4
    assert body["value_columns"] == ["B"]


def test_analyze_bad_extension():
    r = client.post("/analyze",
                    files={"file": ("t.txt", b"hello", "text/plain")})
    assert r.status_code == 400


def test_bad_extension():
    r = client.post("/convert",
                    files={"file": ("t.txt", b"hello", "text/plain")},
                    data={"columns": "C"})
    assert r.status_code == 400


def test_bad_style():
    r = post_convert([[0, "x", 1]], {"columns": "C", "style": "middle"})
    assert r.status_code == 400


def test_auto_columns_when_columns_omitted():
    rows = [["Ур.", "Наим.", "Сумма"],
            [0, "Группа A", None],
            [1, "x", 10],
            [1, "y", 20],
            [0, "Итого A", 30]]
    r = post_convert(rows, {"level_column": "A"})
    assert r.status_code == 200, r.text
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws["C5"].value == "=SUM(C3:C4)"
    assert ws["A5"].value == 0
    assert ws["B5"].value == "Итого A"


def test_unknown_sheet():
    r = post_convert([[0, "x", 1]], {"columns": "C", "sheet": "НетТакого"})
    assert r.status_code == 400
