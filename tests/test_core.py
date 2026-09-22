import openpyxl
import pytest
from openpyxl import Workbook

from subtotals import analyze, process


def write_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def convert(tmp_path, rows, **kwargs):
    src = write_xlsx(tmp_path / "in.xlsx", rows)
    dst = tmp_path / "out.xlsx"
    kwargs.setdefault("level_column", 1)
    process(str(src), str(dst), kwargs.pop("columns", [3]), **kwargs)
    return openpyxl.load_workbook(dst).active


def test_t1_bottom_two_levels(tmp_path):
    ws = convert(tmp_path, [
        [0, "Группа A", None],
        [1, "x", 10],
        [1, "y", 20],
        [0, "Итого A", 30],
    ])
    assert ws["C4"].value == "=SUM(C2:C3)"
    assert ws["C1"].value is None
    assert ws["C2"].value == 10
    assert ws["C3"].value == 20


def test_t2_top_two_levels(tmp_path):
    ws = convert(tmp_path, [
        [0, "Итого A", 30],
        [1, "x", 10],
        [1, "y", 20],
    ])
    assert ws["C1"].value == "=SUM(C2:C3)"


def test_t3_three_levels_reference_subtotals(tmp_path):
    ws = convert(tmp_path, [
        [0, "Итого всё", 0],
        [1, "Итого A", 0],
        [2, "x", 10],
        [2, "y", 20],
        [1, "Итого B", 0],
        [2, "z", 5],
    ])
    assert ws["C1"].value == "=C2+C5"
    assert ws["C2"].value == "=SUM(C3:C4)"
    assert ws["C5"].value == "=C6"


def test_t4_mixed_style(tmp_path):
    ws = convert(tmp_path, [
        [0, "Итого A", 30],
        [1, "x", 10],
        [1, "y", 20],
        [0, "Группа B", None],
        [1, "z", 5],
        [0, "Итого B", 5],
    ])
    assert ws["C1"].value == "=SUM(C2:C3)"
    assert ws["C4"].value is None
    assert ws["C6"].value == "=C5"


def test_root_grand_total_references_group_subtotals(tmp_path):
    ws = convert(tmp_path, [
        [0, "Группа A", None],
        [1, "x", 10],
        [1, "y", 20],
        [0, "Итого A", 30],
        [0, "Группа B", None],
        [1, "z", 5],
        [0, "Итого B", 5],
        [0, "Итого", 35],
    ])
    assert ws["C4"].value == "=SUM(C2:C3)"
    assert ws["C7"].value == "=C6"
    assert ws["C8"].value == "=C4+C7"


def test_t5_subgroup_and_direct_details(tmp_path):
    ws = convert(tmp_path, [
        [0, "Итого A", 0],
        [1, "деталь x", 10],
        [1, "Итого A1", 0],
        [2, "деталь y", 5],
        [1, "деталь z", 7],
    ])
    assert ws["C1"].value == "=C2+C3+C5"
    assert ws["C3"].value == "=C4"
    assert ws["C2"].value == 10
    assert ws["C4"].value == 5
    assert ws["C5"].value == 7


def test_t6_single_child(tmp_path):
    ws = convert(tmp_path, [
        [0, "Итого A", 30],
        [1, "x", 10],
    ])
    assert ws["C1"].value == "=C2"


def test_nested_bottom_groups(tmp_path):
    ws = convert(tmp_path, [
        [0, "Группа A", None],
        [1, "Подгруппа A1", None],
        [2, "x", 10],
        [2, "y", 20],
        [1, "Итого A1", 30],
        [1, "z", 5],
        [0, "Итого A", 35],
    ])
    assert ws["C1"].value is None
    assert ws["C2"].value is None
    assert ws["C5"].value == "=SUM(C3:C4)"
    assert ws["C7"].value == "=SUM(C5:C6)"


def test_forced_style_top(tmp_path):
    ws = convert(tmp_path, [
        [0, "Группа A", None],
        [1, "x", 10],
        [1, "y", 20],
        [0, "Итого A", 30],
    ], forced_style="top")
    assert ws["C1"].value == "=SUM(C2:C3)"
    assert ws["C4"].value == 30


def test_forced_style_bottom(tmp_path):
    ws = convert(tmp_path, [
        [0, "Группа A", None],
        [1, "x", 10],
        [1, "y", 20],
        [0, "Итого A", 30],
    ], forced_style="bottom")
    assert ws["C4"].value == "=SUM(C2:C3)"
    assert ws["C1"].value is None


def test_dry_run_creates_no_file_and_keeps_source(tmp_path):
    src = write_xlsx(tmp_path / "in.xlsx", [
        [0, "Итого A", 30],
        [1, "x", 10],
    ])
    dst = tmp_path / "out.xlsx"
    stats = process(str(src), str(dst), [3], level_column=1, dry_run=True)
    assert not dst.exists()
    assert stats == {"top": 1, "bottom": 0, "cells": 1}
    assert openpyxl.load_workbook(src).active["C1"].value == 30


def test_auto_columns_detect_numeric_totals(tmp_path):
    src = write_xlsx(tmp_path / "in.xlsx", [
        [1, 101, "x", 10, 100],
        [1, 102, "y", 20, 200],
        [0, None, "Итого", 30, 300],
    ])
    dst = tmp_path / "out.xlsx"
    process(str(src), str(dst), level_column=1)
    ws = openpyxl.load_workbook(dst).active
    assert ws["D3"].value == "=SUM(D1:D2)"
    assert ws["E3"].value == "=SUM(E1:E2)"
    assert ws["A3"].value == 0
    assert ws["B3"].value is None


def test_explicit_columns_override_auto(tmp_path):
    ws = convert(tmp_path, [
        [0, "Группа A", None],
        [1, "x", 10],
        [1, "y", 20],
        [0, "Итого A", 30],
    ], columns=[3])
    assert ws["C4"].value == "=SUM(C2:C3)"
    assert ws["C1"].value is None


def test_analyze_flat_report(tmp_path):
    src = write_xlsx(tmp_path / "in.xlsx", [
        ["Сотрудник", "Код", "Часы", "Сумма"],
        ["Иванов", 101, 10, 100],
        ["Петров", 102, 20, 200],
        ["Итого", None, 30, 300],
    ])
    result = analyze(str(src))
    assert result["flat"] is True
    assert result["grand_total"] is True
    assert result["total_row"] == 4
    assert result["value_columns"] == ["C", "D"]
    assert result["cells"] == 2
    assert result["sheet"] == "Sheet"


def test_analyze_grouped_report(tmp_path):
    src = write_xlsx(tmp_path / "in.xlsx", [
        [0, "Группа A", None],
        [1, "x", 10],
        [1, "y", 20],
        [0, "Итого A", 30],
    ])
    result = analyze(str(src), level_column=1)
    assert result["flat"] is False
    assert result["grand_total"] is False
    assert result["grouped_rows"] == 1
    assert result["value_columns"] == ["C"]


def test_grand_total_flat_report(tmp_path):
    src = write_xlsx(tmp_path / "in.xlsx", [
        ["Параметры:", None, None],
        ["Сотрудник", "Нач. остаток", "Кон. остаток"],
        ["Иванов", 91.35, 91.35],
        ["Петров", 19.5, 19.5],
        ["Итого", 110.85, 110.85],
    ])
    dst = tmp_path / "out.xlsx"
    stats = process(str(src), str(dst), [2, 3], grand_total=True)
    assert stats == {"top": 0, "bottom": 1, "cells": 2}
    ws = openpyxl.load_workbook(dst).active
    assert ws["B5"].value == "=SUM(B3:B4)"
    assert ws["C5"].value == "=SUM(C3:C4)"
    assert ws["B3"].value == 91.35


def test_grand_total_auto_columns(tmp_path):
    src = write_xlsx(tmp_path / "in.xlsx", [
        ["Сотрудник", "Код", "Часы", "Сумма"],
        ["Иванов", 101, 10, 100],
        ["Петров", 102, 20, 200],
        ["Итого", None, 30, 300],
    ])
    dst = tmp_path / "out.xlsx"
    stats = process(str(src), str(dst), grand_total=True)
    assert stats == {"top": 0, "bottom": 1, "cells": 2}
    ws = openpyxl.load_workbook(dst).active
    assert ws["C4"].value == "=SUM(C2:C3)"
    assert ws["D4"].value == "=SUM(D2:D3)"
    assert ws["B4"].value is None
    assert ws["A4"].value == "Итого"


def test_grand_total_without_numbers_skips(tmp_path):
    src = write_xlsx(tmp_path / "in.xlsx", [
        ["Сотрудник", "Часы"],
        ["Иванов", None],
        ["Итого", None],
    ])
    dst = tmp_path / "out.xlsx"
    stats = process(str(src), str(dst), [2], grand_total=True)
    assert stats["cells"] == 0


def test_missing_sheet_raises(tmp_path):
    src = write_xlsx(tmp_path / "in.xlsx", [[0, "Итого A", 30], [1, "x", 10]])
    with pytest.raises(ValueError):
        process(str(src), str(tmp_path / "out.xlsx"), [3],
                level_column=1, sheet_name="НетТакого")


def test_outline_levels_without_column(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Итого A", 30])
    ws.append(["x", 10])
    ws.append(["y", 20])
    ws.row_dimensions[2].outlineLevel = 1
    ws.row_dimensions[3].outlineLevel = 1
    src = tmp_path / "in.xlsx"
    wb.save(src)
    dst = tmp_path / "out.xlsx"
    process(str(src), str(dst), [2])
    assert openpyxl.load_workbook(dst).active["B1"].value == "=SUM(B2:B3)"
