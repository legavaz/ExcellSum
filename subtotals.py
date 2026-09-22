#!/usr/bin/env python3
"""Ядро конвертации итогов 1С в формулы Excel."""

import argparse
import sys

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter


class SheetNotFound(ValueError):
    """Запрошенный лист отсутствует в книге."""


def parse_columns(s):
    out = []
    for p in s.split(","):
        p = p.strip()
        if p:
            out.append(int(p) if p.isdigit() else column_index_from_string(p))
    return out


def get_levels(ws, level_column=None):
    levels = {}
    for r in range(1, ws.max_row + 1):
        if level_column:
            v = ws.cell(row=r, column=level_column).value
            try:
                levels[r] = int(v) if v is not None else 0
            except (ValueError, TypeError):
                levels[r] = 0
            continue
        rd = ws.row_dimensions.get(r)
        ol = getattr(rd, "outlineLevel", None) if rd else None
        if ol:
            levels[r] = ol
            continue
        ind = 0
        for c in range(1, min(ws.max_column, 5) + 1):
            cell = ws.cell(row=r, column=c)
            if cell.value is not None and cell.alignment and cell.alignment.indent:
                ind = cell.alignment.indent
                break
        levels[r] = ind
    return levels


def classify(levels, r, max_row, min_row=1, forced=None):
    if forced in ("top", "bottom"):
        return forced if _has_children(levels, r, max_row, min_row, forced) else None
    L = levels.get(r, 0)
    above = levels.get(r - 1, L) if r - 1 >= min_row else L
    below = levels.get(r + 1, L) if r + 1 <= max_row else L
    a_deep, b_deep = above > L, below > L
    if b_deep and not a_deep:
        return "top"
    if a_deep and not b_deep:
        return "bottom"
    return None


def _has_children(levels, r, max_row, min_row, style):
    L = levels.get(r, 0)
    if style == "top":
        return r + 1 <= max_row and levels.get(r + 1, 0) > L
    return r - 1 >= min_row and levels.get(r - 1, 0) > L


def direct_children(levels, r, max_row, style, min_row=1):
    L = levels[r]
    children = []
    if style == "top":
        i = r + 1
        while i <= max_row and levels.get(i, 0) > L:
            if levels[i] == L + 1:
                children.append(i)
                j = i + 1
                while j <= max_row and levels.get(j, 0) > L + 1:
                    j += 1
                i = j
            else:
                i += 1
    else:
        i = r - 1
        while i >= min_row and levels.get(i, 0) > L:
            if levels[i] == L + 1:
                children.append(i)
                j = i - 1
                while j >= min_row and levels.get(j, 0) > L + 1:
                    j -= 1
                i = j
            else:
                i -= 1
        children.sort()
    return children


def build_formula(letter, rows):
    if not rows:
        return None
    if len(rows) == 1:
        return f"={letter}{rows[0]}"
    if rows == list(range(rows[0], rows[-1] + 1)):
        return f"=SUM({letter}{rows[0]}:{letter}{rows[-1]})"
    return "=" + "+".join(f"{letter}{i}" for i in rows)


def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def find_grand_total_row(ws, col_idx, max_row, min_row=1):
    """Последняя строка с числом в колонках значений — кандидат в общий итог."""
    last = None
    for r in range(min_row, max_row + 1):
        if any(_is_number(ws.cell(row=r, column=c).value) for c in col_idx):
            last = r
    return last


def grand_total_plan(ws, col_idx, max_row, min_row=1):
    """Плоский отчёт: последняя числовая строка суммирует непрерывный блок выше."""
    cols = list(col_idx) if col_idx else list(range(1, ws.max_column + 1))
    total = find_grand_total_row(ws, cols, max_row, min_row)
    if total is None:
        return None
    kids = []
    r = total - 1
    while r >= min_row:
        if not any(_is_number(ws.cell(row=r, column=c).value) for c in cols):
            break
        kids.append(r)
        r -= 1
    if not kids:
        return None
    kids.reverse()
    return total, "bottom", kids


def target_columns(ws, total_row, kids, col_idx, exclude=None):
    """Колонки для замены: заданные явно либо авто — числовой итог + числа у детей."""
    if col_idx:
        return col_idx
    cols = []
    for c in range(1, ws.max_column + 1):
        if c == exclude:
            continue
        if not _is_number(ws.cell(row=total_row, column=c).value):
            continue
        if any(_is_number(ws.cell(row=r, column=c).value) for r in kids):
            cols.append(c)
    return cols


class _Reps:
    """Представители поддеревьев: строка с числом либо итог вложенной группы."""

    def __init__(self, ws, children, paired_totals):
        self.ws = ws
        self.children = children
        self.paired_totals = paired_totals
        self.cache = {}

    def rows(self, node, col):
        key = (node, col)
        if key in self.cache:
            return self.cache[key]
        if _is_number(self.ws.cell(row=node, column=col).value):
            out = (node,)
        elif node in self.paired_totals:
            out = self.rows(self.paired_totals[node], col)
        else:
            acc = []
            for child in self.children.get(node, []):
                acc.extend(self.rows(child, col))
            out = tuple(dict.fromkeys(acc))
        self.cache[key] = out
        return out

    def rows_for(self, nodes, col):
        acc = []
        for node in nodes:
            acc.extend(self.rows(node, col))
        return list(dict.fromkeys(acc))


def _block_groups(levels, children, openers_set, r, min_row, block_totals):
    """Группы верхнего уровня в блоке перед итоговой строкой r."""
    L = levels.get(r, 0)
    groups = []
    i = r - 1
    while i >= min_row:
        li = levels.get(i, 0)
        if li < L:
            break
        if li == L and i in openers_set:
            groups.append(i)
        elif (li <= L and not children.get(i)
              and i - 1 >= min_row and levels.get(i - 1, 0) > li
              and i in block_totals):
            break
        i -= 1
    groups.reverse()
    return groups


def build_tree(levels, max_row, min_row=1):
    """Дерево группировок: родитель — ближайшая строка выше с меньшим уровнем."""
    children = {r: [] for r in range(min_row, max_row + 1)}
    subtree_end = {}
    stack = []
    for r in range(min_row, max_row + 1):
        L = levels.get(r, 0)
        while stack and levels.get(stack[-1], 0) >= L:
            subtree_end[stack.pop()] = r - 1
        if stack:
            children[stack[-1]].append(r)
        stack.append(r)
    for r in stack:
        subtree_end[r] = max_row
    return children, subtree_end


def plan_auto(levels, max_row, ws, min_row=1, exclude=None):
    """План замены для style=auto: (план, представители поддеревьев)."""
    children, subtree_end = build_tree(levels, max_row, min_row)
    parent = {r: None for r in range(min_row, max_row + 1)}
    for p, kids in children.items():
        for k in kids:
            parent[k] = p

    openers = [r for r in range(min_row, max_row + 1)
               if children[r] and children[r][0] == r + 1]
    openers_set = set(openers)
    strong_top = [
        r for r in openers
        if not (r - 1 >= min_row and levels.get(r - 1, 0) > levels.get(r, 0))
    ]
    strong_bottom_base = [
        r for r in range(min_row, max_row + 1)
        if r - 1 >= min_row
        and levels.get(r - 1, 0) > levels.get(r, 0)
        and not children[r]
        and (r + 1 > max_row or levels.get(r + 1, 0) <= levels.get(r, 0))
    ]
    guard_inner = len(strong_bottom_base) >= len(strong_top)
    strong_bottom_set = set(strong_bottom_base)

    paired = {}
    bottom_kids = {}
    block_totals = set()
    totals = set()
    plan = []
    reps = _Reps(ws, children, paired)

    def hits(nodes, total_row, cols):
        if not nodes:
            return 0
        found = 0
        for c in cols:
            rows = reps.rows_for(nodes, c)
            if not rows:
                continue
            s = sum(ws.cell(row=i, column=c).value for i in rows)
            if abs(s - ws.cell(row=total_row, column=c).value) < 1e-6:
                found += 1
        return found

    def local_scope(r, L):
        h = r - 1
        while h >= min_row and levels.get(h, 0) > L:
            h -= 1
        if (h >= min_row and h in openers_set and levels.get(h, 0) == L
                and subtree_end.get(h) == r - 1):
            return h, [i for i in children[h] if levels.get(i, 0) == L + 1]
        return None, []

    def choose_root_scope(r, L):
        nonlocal reps
        local_header, local_nodes = local_scope(r, L)
        block_nodes = _block_groups(levels, children, openers_set, r, min_row,
                                    block_totals)
        if not block_nodes:
            block_nodes = [i for i in range(min_row, r)
                           if levels.get(i, 0) == L + 1]
        numeric_cols = [c for c in range(1, ws.max_column + 1)
                        if c != exclude
                        and _is_number(ws.cell(row=r, column=c).value)]
        if local_nodes:
            local_hits = hits(local_nodes, r, numeric_cols)
            block_hits = hits(block_nodes, r, numeric_cols)
            if block_hits > local_hits:
                block_totals.add(r)
                return block_nodes
            if local_header is not None:
                paired[local_header] = r
                reps = _Reps(ws, children, paired)
            return local_nodes
        if block_nodes:
            block_totals.add(r)
        return block_nodes

    for r in range(min_row, max_row + 1):
        is_base = r in strong_bottom_set
        is_cascade = (
            not is_base
            and not children[r]
            and r - 1 >= min_row
            and levels.get(r, 0) == levels.get(r - 1, 0)
            and (r - 1) in totals
            and any(_is_number(ws.cell(row=r, column=c).value)
                    for c in range(1, ws.max_column + 1))
        )
        if not (is_base or is_cascade):
            continue
        L = levels.get(r, 0)
        if parent.get(r) is not None:
            if not guard_inner:
                continue
            h = r - 1
            while h >= min_row and levels.get(h, 0) > L:
                h -= 1
            if (h >= min_row and h in openers_set and levels.get(h, 0) == L
                    and subtree_end.get(h) == r - 1):
                paired[h] = r
                bottom_kids[r] = [
                    i for i in range(h + 1, r)
                    if levels.get(i, 0) == L + 1
                ]
                totals.add(r)
            continue
        nodes = choose_root_scope(r, L)
        if nodes:
            plan.append((r, "bottom", nodes))
            totals.add(r)

    for r in openers:
        if r in paired:
            continue
        L = levels.get(r, 0)
        kids = [i for i in children[r] if levels.get(i, 0) == L + 1]
        if kids:
            plan.append((r, "top", kids))
    for r, kids in bottom_kids.items():
        if kids:
            plan.append((r, "bottom", kids))
    plan.sort(key=lambda item: item[0])
    return plan, reps


def plan_cells(ws, reps, plan, col_idx, exclude=None):
    """Ячейки к замене: (строка, стиль, колонка, строки для формулы)."""
    cells = []
    for r, style, nodes in plan:
        for c in (col_idx or range(1, ws.max_column + 1)):
            if not col_idx and c == exclude:
                continue
            has_value = _is_number(ws.cell(row=r, column=c).value)
            if not has_value and (style == "top" or not col_idx):
                continue
            rows = reps.rows_for(nodes, c)
            if rows:
                cells.append((r, style, c, rows))
    return cells


def analyze(source, sheet_name=None, level_column=None):
    """Разбор файла: группировки, подходящий режим и колонки значений."""
    wb = openpyxl.load_workbook(source)
    if sheet_name is not None:
        if sheet_name not in wb.sheetnames:
            raise SheetNotFound(
                f"Лист {sheet_name!r} не найден. Доступные: {wb.sheetnames}"
            )
        ws = wb[sheet_name]
    else:
        ws = wb.active
    max_row = ws.max_row
    levels = get_levels(ws, level_column)
    plan, reps = plan_auto(levels, max_row, ws, exclude=level_column)
    cells = plan_cells(ws, reps, plan, [], level_column)
    auto_cols = sorted({c for _, _, c, _ in cells})
    auto_rows = len({r for r, _, _, _ in cells})
    grand_plan = grand_total_plan(ws, [], max_row)
    grand_cols = []
    if grand_plan:
        total, _, nodes = grand_plan
        for c in range(1, ws.max_column + 1):
            if c == level_column:
                continue
            if not _is_number(ws.cell(row=total, column=c).value):
                continue
            if any(_is_number(ws.cell(row=i, column=c).value) for i in nodes):
                grand_cols.append(c)
    use_grand = not auto_cols and bool(grand_cols)
    cols = grand_cols if use_grand else auto_cols
    return {
        "sheet": ws.title,
        "sheets": list(wb.sheetnames),
        "rows": max_row,
        "columns": ws.max_column,
        "flat": use_grand,
        "grand_total": use_grand,
        "total_row": grand_plan[0] if use_grand else None,
        "value_columns": [get_column_letter(c) for c in cols],
        "cells": len(grand_cols) if use_grand else len(cells),
        "grouped_rows": 0 if use_grand else auto_rows,
    }


def process(input_path, output_path, value_columns=None, level_column=None,
            sheet_name=None, forced_style=None, dry_run=False,
            grand_total=False):
    wb = openpyxl.load_workbook(input_path)
    if sheet_name is not None:
        if sheet_name not in wb.sheetnames:
            raise SheetNotFound(
                f"Лист {sheet_name!r} не найден. Доступные: {wb.sheetnames}"
            )
        ws = wb[sheet_name]
    else:
        ws = wb.active
    col_idx = [column_index_from_string(c) if isinstance(c, str) else c
               for c in (value_columns or [])]
    levels = get_levels(ws, level_column)
    max_row = ws.max_row
    stats = {"top": 0, "bottom": 0, "cells": 0}

    if grand_total:
        cells = []
        single = grand_total_plan(ws, col_idx, max_row)
        if single:
            r, style, nodes = single
            for c in (col_idx or range(1, ws.max_column + 1)):
                if not col_idx and not _is_number(ws.cell(row=r, column=c).value):
                    continue
                rows = [i for i in nodes
                        if _is_number(ws.cell(row=i, column=c).value)]
                if rows:
                    cells.append((r, style, c, rows))
    elif forced_style is None:
        plan, reps = plan_auto(levels, max_row, ws, exclude=level_column)
        cells = plan_cells(ws, reps, plan, col_idx, level_column)
    else:
        cells = []
        for r in range(1, max_row + 1):
            style = classify(levels, r, max_row, forced=forced_style)
            if style is None:
                continue
            kids = direct_children(levels, r, max_row, style)
            if not kids:
                continue
            for ci in target_columns(ws, r, kids, col_idx, level_column):
                cells.append((r, style, ci, kids))

    styled_rows = {"top": set(), "bottom": set()}
    for r, style, ci, rows in cells:
        styled_rows[style].add(r)
        letter = get_column_letter(ci)
        f = build_formula(letter, rows)
        cell = ws.cell(row=r, column=ci)
        old = cell.value
        if f:
            if not dry_run:
                cell.value = f
            stats["cells"] += 1
            print(f"  r{r:>4} [{style:6}] {letter}: {old!r} -> {f}")

    stats["top"] = len(styled_rows["top"])
    stats["bottom"] = len(styled_rows["bottom"])

    print(f"\nИтогов: сверху {stats['top']}, снизу {stats['bottom']}, "
          f"заменено ячеек {stats['cells']}")
    if not dry_run:
        wb.save(output_path)
        print(f"Сохранено: {output_path}")
    return stats


def main():
    ap = argparse.ArgumentParser(description="1C: итоги → SUM (top/bottom)")
    ap.add_argument("input")
    ap.add_argument("-o", "--output")
    ap.add_argument("-c", "--columns",
                    help="Колонки со значениями, напр. E,F,G (по умолчанию — авто)")
    ap.add_argument("-l", "--level-column")
    ap.add_argument("-s", "--sheet")
    ap.add_argument("--style", choices=["auto", "top", "bottom"], default="auto")
    ap.add_argument("--grand-total", action="store_true",
                    help="Плоский отчёт: последняя числовая строка — общий итог")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    out = a.output or a.input.replace(".xlsx", "_formulas.xlsx")
    if out == a.input and not a.dry_run:
        sys.exit("Ошибка: выходной файл совпадает с входным. Укажите -o.")

    process(a.input, out, parse_columns(a.columns) if a.columns else None,
            column_index_from_string(a.level_column) if a.level_column else None,
            a.sheet, None if a.style == "auto" else a.style, a.dry_run,
            a.grand_total)


if __name__ == "__main__":
    main()
